"""Read a spreadsheet as though it were a paginated document.

Kentucky is the reason this exists. Its scoring criteria are published only as
an .xlsx workbook: the Multifamily Guidelines carry no point values across 107
pages, and the QAP states that KHC makes awards "without determining points".
There is no PDF anywhere to code from.

Rather than teach the whole pipeline a second document model, a workbook is
presented here as a document whose **pages are sheets**. `page_count` is the
number of sheets, `doc[i].get_text()` is sheet i+1 rendered as text, and the
existing citation check -- does this quote appear on the page it cites? --
works unchanged, meaning "does this quote appear on the sheet it cites".

`criteria.cell_ref` then adds the precision a spreadsheet allows and a PDF does
not: `3)Scoring Overview!H7` names one cell. Where it is present the loader
checks the quote against that cell alone, which is a stricter test than any
page citation, because a cell is exact where a page is a haystack.

Cells are laid out with tabs between columns and newlines between rows, so a
row of a scoring table reads as one line and a quote spanning a label and its
value can still be found.
"""

from __future__ import annotations

import re
from pathlib import Path

SHEET_SUFFIXES = {".xlsx", ".xlsm"}
CELL_REF = re.compile(r"^(?P<sheet>.+)!(?P<ref>[A-Z]{1,3}\d{1,7}(?::[A-Z]{1,3}\d{1,7})?)$")


def is_sheet(path: str | Path | None) -> bool:
    return bool(path) and Path(path).suffix.lower() in SHEET_SUFFIXES


def _cells_to_text(rows) -> str:
    out = []
    for row in rows:
        vals = ["" if c.value is None else str(c.value) for c in row]
        while vals and not vals[-1]:
            vals.pop()
        out.append("\t".join(vals))
    while out and not out[-1].strip():
        out.pop()
    return "\n".join(out)


class _Sheet:
    """One worksheet, standing in for a page."""

    def __init__(self, ws):
        self._ws = ws
        self.title = ws.title

    def get_text(self) -> str:
        return _cells_to_text(self._ws.iter_rows())

    def search_for(self, needle: str, **_kw):
        """Present for API compatibility; the loader falls back to text search."""
        return []


class SheetDoc:
    """A workbook wearing just enough of PyMuPDF's Document API."""

    def __init__(self, path: str | Path):
        import openpyxl
        self.path = Path(path)
        # data_only: we want the computed values an applicant would read,
        # not the formulas behind them.
        self._wb = openpyxl.load_workbook(self.path, data_only=True, read_only=False)
        self._sheets = [_Sheet(ws) for ws in self._wb.worksheets]
        self.metadata = {"title": self.path.stem, "producer": "openpyxl"}

    @property
    def page_count(self) -> int:
        return len(self._sheets)

    def __getitem__(self, i: int) -> _Sheet:
        return self._sheets[i]

    def __iter__(self):
        return iter(self._sheets)

    def sheet_names(self) -> list[str]:
        return [s.title for s in self._sheets]

    def cell_text(self, ref: str) -> str | None:
        """Text at `Sheet name!A1` or `Sheet name!A1:C4`, or None if unresolvable."""
        m = CELL_REF.match(ref.strip())
        if not m:
            return None
        name, addr = m.group("sheet"), m.group("ref")
        if name not in self._wb.sheetnames:
            return None
        ws = self._wb[name]
        try:
            got = ws[addr]
        except (ValueError, KeyError):
            return None
        if not isinstance(got, tuple):          # a single cell
            return "" if got.value is None else str(got.value)
        rows = got if isinstance(got[0], tuple) else (got,)
        return _cells_to_text(rows)

    def close(self) -> None:
        self._wb.close()


def open_doc(path: str | Path):
    """Open a PDF or a workbook, whichever this is."""
    if is_sheet(path):
        return SheetDoc(path)
    import pymupdf
    return pymupdf.open(path)
